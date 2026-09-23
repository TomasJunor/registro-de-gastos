import datetime as dt
import shutil
from pathlib import Path

import openpyxl
import pytest
import yaml

from gastos import cli, report
from gastos.categorize import Categorizador, agregar_regla_personal
from gastos.db import DB
from gastos.importers.archivo import importar_archivo
from gastos.importers.mercadopago import parsear_reporte
from gastos.models import Movimiento
from gastos.parsing import parse_fecha, parse_importe

FIX = Path(__file__).parent / "fixtures"
RAIZ = Path(__file__).parent.parent


# --- parsing -------------------------------------------------------------------

@pytest.mark.parametrize("texto,esperado", [
    ("1.234,56", 1234.56),
    ("$ -1.234,56", -1234.56),
    ("(1.234,56)", -1234.56),
    ("1.234,56-", -1234.56),
    ("1234.56", 1234.56),
    ("1.234", 1234.0),
    ("1.234.567", 1234567.0),
    ("12,5", 12.5),
    ("U$S 10,00", 10.0),
    ("-1,234.56", -1234.56),
    (1500, 1500.0),
    ("", None),
])
def test_parse_importe(texto, esperado):
    assert parse_importe(texto) == esperado


def test_parse_importe_punto_decimal():
    assert parse_importe("1234.5", decimal=".") == 1234.5
    assert parse_importe("1,234", decimal=".") == 1234.0


@pytest.mark.parametrize("texto,esperado", [
    ("05/09/2026", dt.date(2026, 9, 5)),
    ("5/9/26", dt.date(2026, 9, 5)),
    ("2026-09-05T10:00:00.000-03:00", dt.date(2026, 9, 5)),
    ("05-Sep-2026", dt.date(2026, 9, 5)),
    ("05 set 26", dt.date(2026, 9, 5)),
    (dt.datetime(2026, 9, 5, 10), dt.date(2026, 9, 5)),
])
def test_parse_fecha(texto, esperado):
    assert parse_fecha(texto) == esperado


# --- importadores -------------------------------------------------------------

def test_importar_csv_santander_con_debito_credito():
    movs = importar_archivo(FIX / "santander_cuenta.csv", {"id": "santander_pesos"})
    montos = [m.monto for m in movs]
    assert montos == [1250000.50, -350000.0, -2100.0, -100000.0, -45321.99, -12500.0, -12500.0]
    assert movs[0].descripcion == "Acreditación de haberes ACME SA"
    assert movs[0].fecha == dt.date(2026, 9, 1)
    assert movs[0].extra.get("Referencia") == "123"


def _visa_xlsx(ruta: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Tarjeta Visa terminada en 1234"])
    ws.append([])
    ws.append(["Fecha", "Descripción", "Cuotas", "Comprobante", "Pesos", "Dólares"])
    ws.append([dt.datetime(2026, 9, 2), "NETFLIX.COM", None, "0001", None, 9.99])
    ws.append(["03/09/2026", "PEDIDOSYA*RESTO", None, "0002", "15.300,00", None])
    ws.append(["04/09/2026", "FRAVEGA SUC 12", "C.01/06", "0003", 80000.0, None])
    ws.append(["06/09/2026", "SU PAGO EN PESOS", None, None, -350000.0, None])
    ws.append([None, "Total", None, None, -254700.0, 9.99])
    wb.save(ruta)


def test_importar_xlsx_visa_con_columnas_por_moneda(tmp_path):
    ruta = tmp_path / "visa.xlsx"
    _visa_xlsx(ruta)
    movs = importar_archivo(ruta, {"id": "visa", "invertir_signo": True})
    assert [(m.descripcion, m.moneda, m.monto) for m in movs] == [
        ("NETFLIX.COM", "USD", -9.99),
        ("PEDIDOSYA*RESTO", "ARS", -15300.0),
        ("FRAVEGA SUC 12", "ARS", -80000.0),
        ("SU PAGO EN PESOS", "ARS", 350000.0),
    ]
    assert movs[2].extra["Cuotas"] == "C.01/06"


def test_importar_xls_que_en_realidad_es_html(tmp_path):
    ruta = tmp_path / "movimientos.xls"
    ruta.write_text(
        "<html><body><table>"
        "<tr><td>Fecha</td><td>Concepto</td><td>Importe</td></tr>"
        "<tr><td>10/09/2026</td><td>SPOTIFY</td><td>-4.599,00</td></tr>"
        "<tr><td>11/09/2026</td><td>REINTEGRO PROMO</td><td>1.000,00</td></tr>"
        "</table></body></html>", encoding="utf-8")
    movs = importar_archivo(ruta, {"id": "x"})
    assert [(m.descripcion, m.monto) for m in movs] == [("SPOTIFY", -4599.0), ("REINTEGRO PROMO", 1000.0)]


def test_parsear_reporte_mercadopago():
    texto = (FIX / "mercadopago_release.csv").read_text()
    movs = parsear_reporte(texto, "mercadopago", detalles={"444": {"description": "Netflix"}})
    assert [m.monto for m in movs] == [100000.0, -8500.0, 1234.56, -20000.0]
    assert movs[1].descripcion == "Pago - PEDIDOSYA - [account_money]"
    assert movs[2].descripcion == "Rendimientos"
    assert movs[3].descripcion == "Pago - Netflix - [account_money]"
    assert len({m.id_externo for m in movs}) == 4


# --- categorización -----------------------------------------------------------

@pytest.fixture
def categorizador():
    return Categorizador.desde_archivos(RAIZ / "categorias.yaml")


@pytest.mark.parametrize("desc,monto,esperado", [
    ("Acreditación de haberes ACME SA", 1000, "Sueldo"),
    ("Pago tarjeta Visa", -1000, "Pago de tarjeta"),
    ("SU PAGO EN PESOS", 1000, "Pago de tarjeta"),
    ("Impuesto Ley 25.413 débito", -10, "Impuestos y retenciones"),
    ("Transf. inmediata a Mercado Pago CVU", -1000, "Transferencia entre cuentas propias"),
    ("Débito automático EDENOR", -1000, "Servicios del hogar"),
    ("Compra con débito COTO SUC 45", -1000, "Supermercado"),
    ("PEDIDOSYA*RESTO", -1000, "Delivery y restaurantes"),
    ("NETFLIX.COM", -10, "Suscripciones digitales"),
    ("Rendimientos", 10, "Rendimientos e intereses"),
    ("TRANSFERENCIA DE JUAN PEREZ", 5000, "Transferencias recibidas"),
    ("TRANSFERENCIA A JUAN PEREZ", -5000, "Transferencias enviadas"),
    ("PRESTAMO PERSONAL CUOTA 3", -5000, "Préstamos (cuotas)"),
    ("ALGO QUE NO EXISTE", -1, None),
])
def test_reglas(categorizador, desc, monto, esperado):
    assert categorizador.categorizar(desc, "cualquiera", monto) == esperado


def test_reglas_personales_tienen_prioridad(tmp_path):
    personal = tmp_path / "personales.yaml"
    agregar_regla_personal(personal, "Mascotas", "egreso", "COTO")
    agregar_regla_personal(personal, "Mascotas", "egreso", "PUPPIS")
    cat = Categorizador.desde_archivos(personal, RAIZ / "categorias.yaml")
    assert cat.categorizar("COTO SUC 45", "x", -1) == "Mascotas"
    assert yaml.safe_load(personal.read_text())["categorias"][0]["patrones"] == ["COTO", "PUPPIS"]


# --- base de datos -------------------------------------------------------------

def test_dedup_respeta_movimientos_identicos(tmp_path):
    db = DB(tmp_path / "g.db")
    m = Movimiento("c", dt.date(2026, 9, 5), "CAFE", -1500)
    assert db.insertar([m, m]) == 2          # dos cafés iguales el mismo día
    assert db.insertar([m, m]) == 0          # reimportar el mismo archivo no duplica
    assert db.insertar([m, m, m]) == 1       # un archivo más completo agrega solo el nuevo
    assert len(db.consultar()) == 3


# --- de punta a punta ---------------------------------------------------------

def _config(tmp_path: Path) -> Path:
    shutil.copy(RAIZ / "categorias.yaml", tmp_path / "categorias.yaml")
    cfg = {
        "cuentas": [
            {"id": "santander_pesos", "importador": "archivo", "carpeta": "bandeja/santander"},
            {"id": "visa_santander", "importador": "archivo", "carpeta": "bandeja/visa",
             "invertir_signo": True},
        ]
    }
    ruta = tmp_path / "config.yaml"
    ruta.write_text(yaml.safe_dump(cfg))
    (tmp_path / "bandeja/santander").mkdir(parents=True)
    (tmp_path / "bandeja/visa").mkdir(parents=True)
    shutil.copy(FIX / "santander_cuenta.csv", tmp_path / "bandeja/santander/ext.csv")
    _visa_xlsx(tmp_path / "bandeja/visa/visa.xlsx")
    return ruta


def test_sync_resumen_y_excel(tmp_path, capsys):
    config = _config(tmp_path)
    with pytest.raises(SystemExit) as fin:
        cli.main(["-c", str(config), "sync"])
    assert fin.value.code == 0
    salida = capsys.readouterr().out
    assert "santander_pesos: 7 movimientos nuevos" in salida
    assert "visa_santander: 4 movimientos nuevos" in salida
    # los archivos quedan archivados y la bandeja vacía
    assert not list((tmp_path / "bandeja/santander").iterdir())
    assert (tmp_path / "data/procesados/santander_pesos/ext.csv").exists()

    db = DB(tmp_path / "data/gastos.db")
    cat = Categorizador.desde_archivos(tmp_path / "categorias.yaml")
    df = report.dataframe(db, cat)
    tabla = report.resumen_mensual(df, "ARS")
    sep = "2026-09"
    assert tabla.loc[("TOTAL", "Ingresos"), sep] == pytest.approx(1250000.50)
    # egresos: impuesto 2100 + edenor 45321.99 + coto 25000 + pedidosya 15300 + fravega 80000
    assert tabla.loc[("TOTAL", "Egresos"), sep] == pytest.approx(167721.99)
    # pago de tarjeta (ambos lados) y transferencia a MP no cuentan
    assert ("Egresos", "Pago de tarjeta") not in tabla.index
    assert report.resumen_mensual(df, "USD").loc[("TOTAL", "Egresos"), sep] == pytest.approx(9.99)

    with pytest.raises(SystemExit):
        cli.main(["-c", str(config), "exportar", "--salida", "reporte.xlsx"])
    libro = openpyxl.load_workbook(tmp_path / "reporte.xlsx")
    assert {"Resumen ARS", "Resumen USD", "Movimientos", "Neto por cuenta"} <= set(libro.sheetnames)
    db.close()


def test_asignar_manual_no_se_pisa_al_recategorizar(tmp_path, capsys):
    config = _config(tmp_path)
    with pytest.raises(SystemExit):
        cli.main(["-c", str(config), "sync"])
    db = DB(tmp_path / "data/gastos.db")
    fravega = [f for f in db.consultar() if "FRAVEGA" in f["descripcion"]][0]
    with pytest.raises(SystemExit):
        cli.main(["-c", str(config), "asignar", fravega["id"][:10], "Regalos"])
    with pytest.raises(SystemExit):
        cli.main(["-c", str(config), "recategorizar"])
    fila = db.con.execute("SELECT categoria FROM movimientos WHERE id = ?", (fravega["id"],)).fetchone()
    assert fila["categoria"] == "Regalos"
    db.close()
